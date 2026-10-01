# pyrefly: ignore [missing-import]
from pptx import Presentation
# pyrefly: ignore [missing-import]
from pptx.util import Inches, Pt
# pyrefly: ignore [missing-import]
from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()

    # 1. Title Slide
    slide_layout = prs.slide_layouts[0] 
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Agentic Maritime Brokerage Platform"
    subtitle.text = "AI-Powered Route Intelligence & Quotations\nInfosys Project Presentation"

    # 2. Problem Statement
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "The Problem"
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "Traditional maritime freight brokerage faces several challenges:"
    p = tf.add_paragraph()
    p.text = "Manual and time-consuming quotation processes"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Inability to dynamically weigh multiple routing factors at scale"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Lack of transparency in why specific routes are recommended"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Static pricing that fails to account for real-time risk (weather, transshipments)"
    p.level = 1

    # 3. Our Solution
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "Our Solution"
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "A Multi-Agent AI Platform that automates and optimizes maritime logistics:"
    p = tf.add_paragraph()
    p.text = "Agentic Architecture: Specialized AI agents handle different aspects of the quotation."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Dynamic Scoring: Mathematical models balance Cost, Time, and Risk."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Explainable AI: Generates human-readable reasoning for every recommendation."
    p.level = 1

    # 4. Technical Architecture
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "Technical Architecture"
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "Built using a modern, scalable tech stack:"
    p = tf.add_paragraph()
    p.text = "Frontend: React & Vite (Responsive, modern web UI with custom Canvas charting)"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Backend API: FastAPI (Python-based, high-performance async API)"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "AI Logic: Python Pandas for data processing & custom heuristic agents"
    p.level = 1

    # 5. Milestone 1: Route Intelligence
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "Milestone 1: Route Intelligence"
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "The 'Route Agent' is currently live and evaluates routes using a weighted algorithm:"
    p = tf.add_paragraph()
    p.text = "35% - Transit Time (Prioritizing faster delivery)"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "30% - Cost (Minimizing per-container freight expenses)"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "20% - Transshipments (Reducing cargo handling risk)"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "15% - Distance (Optimizing fuel and nautical miles)"
    p.level = 1

    # 6. Future Roadmap
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "Future Roadmap (Milestone 2+)"
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "The RootAgent orchestrator is ready to integrate new specialized agents:"
    p = tf.add_paragraph()
    p.text = "Pricing Agent: Dynamic fuel surcharges and cargo-type multipliers."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Weather Agent: Real-time risk assessment impacting transit days."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Customs Agent: Automated compliance and tariff estimations."
    p.level = 1

    # 7. Conclusion
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Thank You"
    subtitle.text = "Q&A Session\n\n(Live Demonstration Available)"

    prs.save("AI_Maritime_Brokerage_Presentation.pptx")
    print("Presentation generated successfully!")

if __name__ == '__main__':
    create_presentation()
